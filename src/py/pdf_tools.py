"""Local PDF worker for AnEdiKit.

Uses pypdf for structural PDF work and PyMuPDF for rendering, redaction and
compression.  Optional integrations (LibreOffice and Tesseract) report a
clear error when they are not installed instead of sending a document away.
"""
import argparse
import base64
import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

_weasyprint_dll_handles = []


def configure_weasyprint_windows():
    """Make common GTK3/MSYS2 runtimes visible before importing WeasyPrint."""
    if os.name != "nt":
        return []
    candidates = []
    configured = os.environ.get("WEASYPRINT_DLL_DIRECTORIES", "")
    candidates.extend(Path(item) for item in configured.split(os.pathsep) if item)
    candidates.extend([
        Path(r"C:\GTK3\bin"),
        Path(r"C:\Program Files\GTK3-Runtime Win64\bin"),
        Path(r"C:\Program Files\GTK3-Runtime Win32\bin"),
        Path(r"C:\msys64\ucrt64\bin"),
        Path(r"C:\msys64\mingw64\bin"),
        Path(r"C:\msys64\mingw32\bin"),
    ])
    valid = []
    for folder in candidates:
        if folder.is_dir() and (folder / "libgobject-2.0-0.dll").exists():
            valid.append(str(folder))
            if hasattr(os, "add_dll_directory"):
                _weasyprint_dll_handles.append(os.add_dll_directory(str(folder)))
    if valid:
        os.environ["WEASYPRINT_DLL_DIRECTORIES"] = os.pathsep.join(dict.fromkeys(valid))
        os.environ["PATH"] = os.pathsep.join(dict.fromkeys(valid + [os.environ.get("PATH", "")]))
    return valid


def progress(pct, message):
    print(json.dumps({"type": "progress", "pct": pct, "msg": message}, separators=(",", ":")), flush=True)


def require_pdf_libs():
    try:
        import fitz
        from pypdf import PdfReader, PdfWriter
        return fitz, PdfReader, PdfWriter
    except ImportError as exc:
        raise RuntimeError("PDF support is not installed. Run: pip install -r requirements.txt") from exc


def pages_from_spec(spec, count):
    if not spec.strip():
        return list(range(count))
    pages = []
    for part in spec.replace(" ", "").split(","):
        if not part:
            continue
        bounds = part.split("-", 1)
        try:
            start = max(1, int(bounds[0]))
            end = max(start, int(bounds[-1]))
        except ValueError as exc:
            raise ValueError("Use page numbers such as 1-3, 5, 8-10") from exc
        pages.extend(range(start - 1, min(end, count)))
    return list(dict.fromkeys(pages))


def output_path(params, suffix, extension=".pdf"):
    custom = params.get("output_path", "").strip()
    source = Path(params["files"][0])
    if custom:
        candidate = Path(custom)
        if candidate.suffix:
            candidate.parent.mkdir(parents=True, exist_ok=True)
            return candidate
        candidate.mkdir(parents=True, exist_ok=True)
        return candidate / f"{source.stem}{suffix}{extension}"
    return source.with_name(f"{source.stem}{suffix}{extension}")


def write_selected(reader, indices, destination):
    _, _, PdfWriter = require_pdf_libs()
    writer = PdfWriter()
    for index in indices:
        writer.add_page(reader.pages[index])
    with open(destination, "wb") as stream:
        writer.write(stream)


def office_convert(source, destination, target_format):
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if not office:
        raise RuntimeError("LibreOffice is required for this conversion. Install it, then retry.")
    destination.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run([office, "--headless", "--convert-to", target_format, "--outdir", str(destination.parent), str(source)], check=True)
    generated = destination.parent / f"{source.stem}.{target_format.split(':')[0]}"
    if generated != destination and generated.exists():
        generated.replace(destination)


def extract_text(path):
    fitz, _, _ = require_pdf_libs()
    document = fitz.open(path)
    return "\n\n".join(page.get_text("text") for page in document)


def stamp_pdf(source, destination, label, numbers=False, position="center"):
    from io import BytesIO
    from reportlab.pdfgen import canvas
    _, PdfReader, PdfWriter = require_pdf_libs()
    reader, writer = PdfReader(str(source)), PdfWriter()
    for index, page in enumerate(reader.pages, 1):
        width, height = float(page.mediabox.width), float(page.mediabox.height)
        buffer = BytesIO(); layer = canvas.Canvas(buffer, pagesize=(width, height))
        layer.setFillColorRGB(0.35, 0.35, 0.35); layer.setFont("Helvetica", 9)
        if numbers: layer.drawCentredString(width / 2, 16, str(index))
        else:
            text_width = layer.stringWidth(label, "Helvetica", 9)
            x = width / 2; y = height / 2
            if position == "bottom-right": x, y = width - text_width / 2 - 18, 18
            elif position == "bottom-center": x, y = width / 2, 18
            elif position == "top-right": x, y = width - text_width / 2 - 18, height - 24
            layer.drawCentredString(x, y, label)
        layer.save(); buffer.seek(0)
        page.merge_page(PdfReader(buffer).pages[0]); writer.add_page(page)
    with open(destination, "wb") as stream: writer.write(stream)


def run(params):
    tool = params["tool"]
    files = [Path(value) for value in params.get("files", [])]
    if tool != "HTML to PDF" and not files:
        raise ValueError("Choose at least one source file.")
    progress(5, f"Preparing {tool}")

    if tool == "Merge PDF":
        _, PdfReader, PdfWriter = require_pdf_libs(); writer = PdfWriter()
        for index, path in enumerate(files):
            progress(10 + int(index * 70 / len(files)), f"Adding {path.name}")
            writer.append(str(path))
        destination = output_path(params, "_merged")
        with open(destination, "wb") as stream: writer.write(stream)

    elif tool in {"Split PDF", "Extract Pages", "Remove Pages", "Organize PDF", "Rotate PDF", "Crop PDF"}:
        _, PdfReader, _ = require_pdf_libs(); reader = PdfReader(str(files[0]))
        if tool == "Organize PDF" and len(files) > 1:
            # The organize queue can contain whole PDFs; flatten them in the
            # exact drag-and-drop order so users can assemble documents easily.
            from pypdf import PdfWriter
            combined = PdfWriter()
            for path in files:
                source_reader = PdfReader(str(path)); [combined.add_page(page) for page in source_reader.pages]
            temp = output_path(params, "_organize_source")
            with open(temp, "wb") as stream: combined.write(stream)
            reader = PdfReader(str(temp))
        selected = pages_from_spec(params.get("pages", ""), len(reader.pages))
        if tool == "Split PDF":
            folder = output_path(params, "_split", ""); folder.mkdir(parents=True, exist_ok=True)
            if params.get("split_mode", "individual") == "combined":
                destination = folder / f"{files[0].stem}_selected.pdf"; write_selected(reader, selected, destination)
            else:
                for number, page_index in enumerate(selected, 1):
                    write_selected(reader, [page_index], folder / f"{files[0].stem}_page_{page_index + 1}.pdf")
                destination = folder
        elif tool == "Remove Pages":
            destination = output_path(params, "_pages_removed"); write_selected(reader, [i for i in range(len(reader.pages)) if i not in selected], destination)
        else:
            if tool == "Rotate PDF":
                degrees = int(params.get("rotate_degrees", 90));
                for index in selected: reader.pages[index].rotate(degrees)
                destination = output_path(params, "_rotated"); write_selected(reader, range(len(reader.pages)), destination)
            elif tool == "Crop PDF":
                for index in selected:
                    box = reader.pages[index].mediabox
                    margin = float(params.get("crop_margin", 18)); box.lower_left = (box.left + margin, box.bottom + margin); box.upper_right = (box.right - margin, box.top - margin)
                destination = output_path(params, "_cropped"); write_selected(reader, range(len(reader.pages)), destination)
            else:
                destination = output_path(params, "_extracted" if tool == "Extract Pages" else "_organized"); write_selected(reader, selected if tool == "Extract Pages" else (selected or range(len(reader.pages))), destination)

    elif tool in {"JPG to PDF", "Scan to PDF"}:
        from PIL import Image
        from PIL import ImageOps
        images = [Image.open(path).convert("RGB") for path in files]
        if tool == "JPG to PDF":
            margin = max(0, int(float(params.get("page_margin", 18))))
            orientation = params.get("page_orientation", "auto")
            prepared = []
            for image in images:
                width, height = image.size
                landscape = orientation == "landscape" or (orientation == "auto" and width > height)
                page = (max(width, height), min(width, height)) if landscape else (min(width, height), max(width, height))
                canvas = Image.new("RGB", (page[0] + margin * 2, page[1] + margin * 2), "white")
                fitted = ImageOps.contain(image, (page[0], page[1]))
                canvas.paste(fitted, (margin + (page[0] - fitted.width) // 2, margin + (page[1] - fitted.height) // 2))
                prepared.append(canvas)
            images = prepared
        destination = output_path(params, "_images" if tool == "JPG to PDF" else "_scan")
        images[0].save(destination, save_all=True, append_images=images[1:])

    elif tool == "PDF to JPG":
        fitz, _, _ = require_pdf_libs(); document = fitz.open(files[0]); folder = output_path(params, "_images", ""); folder.mkdir(parents=True, exist_ok=True)
        scale = max(0.5, float(params.get("render_dpi", 150)) / 72.0); image_format = params.get("image_format", "jpg").lower();
        for index, page in enumerate(document):
            page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False).save(str(folder / f"page_{index + 1}.{image_format}"))
        destination = folder

    elif tool in {"Compress PDF", "Repair PDF"}:
        fitz, PdfReader, _ = require_pdf_libs(); destination = output_path(params, "_compressed" if tool == "Compress PDF" else "_repaired")
        if tool == "Compress PDF":
            document = fitz.open(files[0]); level = params.get("compression_level", "balanced"); garbage = 4 if level != "quality" else 2; document.save(destination, garbage=garbage, deflate=True, clean=True)
        else:
            reader = PdfReader(str(files[0]), strict=False); write_selected(reader, range(len(reader.pages)), destination)

    elif tool in {"Protect PDF", "Unlock PDF"}:
        _, PdfReader, PdfWriter = require_pdf_libs(); reader = PdfReader(str(files[0])); password = params.get("password", "")
        if tool == "Unlock PDF" and reader.is_encrypted and not reader.decrypt(password): raise RuntimeError("The supplied password could not unlock this PDF.")
        writer = PdfWriter(); [writer.add_page(page) for page in reader.pages]
        destination = output_path(params, "_protected" if tool == "Protect PDF" else "_unlocked")
        if tool == "Protect PDF":
            if not password: raise ValueError("Enter a password to protect the document.")
            writer.encrypt(password)
        with open(destination, "wb") as stream: writer.write(stream)

    elif tool == "Redact PDF":
        fitz, _, _ = require_pdf_libs(); terms = [term.strip() for term in params.get("pages", "").split(",") if term.strip()]
        if not terms: raise ValueError("Enter the text to redact in Page selection.")
        document = fitz.open(files[0])
        for page in document:
            for term in terms:
                for rect in page.search_for(term): page.add_redact_annot(rect, fill=(0, 0, 0))
            page.apply_redactions()
        destination = output_path(params, "_redacted"); document.save(destination)

    elif tool in {"PDF to Markdown", "AI Summarizer", "Translate PDF", "Compare PDF"}:
        if tool == "Compare PDF":
            if len(files) < 2: raise ValueError("Select the original and revised PDF.")
            text = "# PDF comparison\n\n" + "\n".join(__import__("difflib").unified_diff(extract_text(files[0]).splitlines(), extract_text(files[1]).splitlines(), fromfile=files[0].name, tofile=files[1].name))
            destination = output_path(params, "_comparison", ".md")
        else:
            text = extract_text(files[0])
            if tool == "AI Summarizer":
                sentences = re.split(r"(?<=[.!?])\s+", text); text = "# Summary\n\n" + "\n".join(f"- {s}" for s in sentences[:12])
            elif tool == "Translate PDF":
                try:
                    import argostranslate.translate
                    target_codes = {"English": "en", "Spanish": "es", "French": "fr"}
                    target = target_codes.get(params.get("ai_option"), "en")
                    translation = argostranslate.translate.get_translation_from_codes("en", target)
                    if translation is None: raise RuntimeError("Install the required Argos Translate language package, then retry.")
                    text = f"# Translation ({params.get('ai_option')})\n\n" + translation.translate(text)
                except ImportError as exc:
                    raise RuntimeError("Translate PDF requires the optional Argos Translate local language engine.") from exc
            destination = output_path(params, "_summary" if tool == "AI Summarizer" else "_markdown", ".md")
        destination.write_text(text, encoding="utf-8")

    elif tool in {"PDF to Word", "PDF to Excel", "PDF to PowerPoint"}:
        text = extract_text(files[0])
        if tool == "PDF to Word":
            from docx import Document
            document = Document(); [document.add_paragraph(line) for line in text.splitlines() if line.strip()]; destination = output_path(params, "_editable", ".docx"); document.save(destination)
        elif tool == "PDF to Excel":
            from openpyxl import Workbook
            book = Workbook(); sheet = book.active; [sheet.append([line]) for line in text.splitlines() if line.strip()]; destination = output_path(params, "_tables", ".xlsx"); book.save(destination)
        else:
            from pptx import Presentation
            deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[1]); slide.shapes.title.text = files[0].stem; slide.placeholders[1].text = text[:4000]; destination = output_path(params, "_slides", ".pptx"); deck.save(destination)

    elif tool in {"Word to PDF", "PowerPoint to PDF", "Excel to PDF"}:
        destination = output_path(params, "", ".pdf"); office_convert(files[0], destination, "pdf")

    elif tool in {"Page Numbers", "Watermark", "Edit PDF", "Sign PDF"}:
        suffix = {"Page Numbers": "_numbered", "Watermark": "_watermarked", "Edit PDF": "_edited", "Sign PDF": "_signed"}[tool]
        destination = output_path(params, suffix)
        if tool == "Sign PDF":
            label = params.get("signature_text") or "Approved"
            if params.get("signer_name"): label += f" — {params['signer_name']}"
            if params.get("sign_date"): label += f" ({params['sign_date']})"
        else:
            label = params.get("watermark") or ("Signed with AnEdiKit" if tool == "Sign PDF" else "Edited with AnEdiKit")
        stamp_pdf(files[0], destination, label, numbers=tool == "Page Numbers", position=params.get("sign_position", "center"))

    elif tool == "PDF Forms":
        _, PdfReader, PdfWriter = require_pdf_libs(); reader = PdfReader(str(files[0]), strict=False); writer = PdfWriter(); writer.clone_document_from_reader(reader)
        values = {}
        for entry in params.get("form_values", "").split(","):
            if "=" in entry:
                key, value = entry.split("=", 1); values[key.strip()] = value.strip()
        if values:
            for page in writer.pages: writer.update_page_form_field_values(page, values)
        destination = output_path(params, "_forms")
        with open(destination, "wb") as stream: writer.write(stream)

    elif tool == "PDF to PDF/A":
        ghostscript = shutil.which("gswin64c") or shutil.which("gswin32c") or shutil.which("gs")
        if not ghostscript: raise RuntimeError("PDF/A conversion requires Ghostscript. Install it, then retry.")
        destination = output_path(params, "_archive")
        subprocess.run([ghostscript, "-dPDFA=2", "-dBATCH", "-dNOPAUSE", "-sDEVICE=pdfwrite", f"-sOutputFile={destination}", str(files[0])], check=True)

    elif tool == "OCR PDF":
        try:
            import pytesseract
            from PIL import Image
        except ImportError as exc:
            raise RuntimeError("OCR requires pytesseract and the Tesseract desktop engine.") from exc
        fitz, _, _ = require_pdf_libs(); source = fitz.open(files[0]); destination = output_path(params, "_ocr")
        result = fitz.open()
        for index, page in enumerate(source):
            progress(10 + int(index * 80 / max(1, len(source))), f"Recognizing page {index + 1}")
            pixmap = page.get_pixmap(matrix=fitz.Matrix(max(1, float(params.get("ocr_dpi", 300)) / 150), max(1, float(params.get("ocr_dpi", 300)) / 150)), alpha=False)
            image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            ocr_pdf = pytesseract.image_to_pdf_or_hocr(image, extension="pdf", lang=params.get("ocr_language", "eng"))
            ocr_doc = fitz.open("pdf", ocr_pdf); result.insert_pdf(ocr_doc)
        result.save(destination)

    elif tool == "HTML to PDF":
        try:
            configure_weasyprint_windows()
            from weasyprint import HTML
        except ImportError as exc:
            raise RuntimeError("HTML to PDF requires WeasyPrint. Install it with the PDF dependencies.") from exc
        except OSError as exc:
            raise RuntimeError("WeasyPrint is installed, but Windows GTK/Pango DLLs are unavailable. Install the GTK3 runtime, add its bin folder to WEASYPRINT_DLL_DIRECTORIES, then retry. [WEASYPRINT_WINDOWS_SETUP]") from exc
        destination = Path(params.get("output_path") or Path.cwd() / "webpage.pdf"); HTML(url=params["url"]).write_pdf(destination)

    else:
        raise RuntimeError(f"{tool} requires a dedicated interactive editor or optional local engine and is not available in this build yet.")

    progress(100, f"Completed: {destination}")
    return {"success": True, "output_path": str(destination)}


def preview_pages(path):
    if Path(path).suffix.lower() in {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff"}:
        from PIL import Image
        from io import BytesIO
        image = Image.open(path).convert("RGB"); buffer = BytesIO(); image.thumbnail((900, 1200)); image.save(buffer, format="PNG")
        return ["data:image/png;base64," + base64.b64encode(buffer.getvalue()).decode("ascii")]
    fitz, _, _ = require_pdf_libs()
    document = fitz.open(path)
    previews = []
    for page in document:
        pixmap = page.get_pixmap(matrix=fitz.Matrix(0.8, 0.8), alpha=False)
        previews.append("data:image/png;base64," + base64.b64encode(pixmap.tobytes("png")).decode("ascii"))
    return previews


def pdf_info(path):
    source = Path(path)
    result = {"name": source.name, "size_bytes": source.stat().st_size, "pages": 0, "width": None, "height": None}
    if source.suffix.lower() in {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff"}:
        from PIL import Image
        with Image.open(source) as image:
            result.update(pages=1, width=image.width, height=image.height)
        return result
    _, PdfReader, _ = require_pdf_libs()
    reader = PdfReader(str(source), strict=False)
    result["pages"] = len(reader.pages)
    if reader.pages:
        page = reader.pages[0]
        result["width"] = float(page.mediabox.width); result["height"] = float(page.mediabox.height)
    return result


def main():
    parser = argparse.ArgumentParser(); parser.add_argument("--params"); parser.add_argument("--preview"); parser.add_argument("--info"); args = parser.parse_args()
    try:
        if args.info:
            print(json.dumps(pdf_info(args.info)), flush=True)
        elif args.preview:
            print(json.dumps(preview_pages(args.preview)), flush=True)
        else:
            result = run(json.loads(args.params)); print("ANEDIKIT_RESULT:" + json.dumps(result), flush=True)
    except Exception as exc:
        print("ANEDIKIT_RESULT:" + json.dumps({"success": False, "error": str(exc)}), flush=True); raise


if __name__ == "__main__": main()
